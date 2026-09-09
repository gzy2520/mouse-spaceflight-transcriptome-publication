#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate the publication Methods Word document (.docx) organized strictly in analytical and Supplementary Table order."""

import os
import subprocess
import tempfile
import zipfile
import functools
import lxml.etree as etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

@functools.lru_cache(maxsize=None)
def latex_to_omath(latex_expr):
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
        tmp_path = tmp.name
    try:
        subprocess.run(
            ['pandoc', '-f', 'markdown', '-t', 'docx', '-o', tmp_path],
            input=f'${latex_expr}$'.encode('utf-8'),
            check=True
        )
        with zipfile.ZipFile(tmp_path) as z:
            t = etree.fromstring(z.read('word/document.xml'))
            m = t.find('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath')
            return etree.tostring(m, encoding='unicode')
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)

def insert_in_schema_order(parent, element, tag_order):
    """Insert an OOXML child before the first later-schema child."""
    element_name = element.tag.rsplit("}", 1)[-1]
    element_index = tag_order.index(element_name)
    for index, child in enumerate(parent):
        child_name = child.tag.rsplit("}", 1)[-1]
        if child_name in tag_order and tag_order.index(child_name) > element_index:
            parent.insert(index, element)
            return
    parent.append(element)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:fill="{fill_hex}"/>')
    insert_in_schema_order(
        tcPr,
        shd,
        ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange"],
    )

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:start w:w="{left}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:end w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    insert_in_schema_order(
        tcPr,
        tcMar,
        ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange"],
    )

def set_three_line_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="12" w:space="0" w:color="1F4E78"/>'
        f'  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="1F4E78"/>'
        f'  <w:insideH w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    insert_in_schema_order(
        tblPr,
        borders,
        ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription", "tblPrChange"],
    )
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="2E75B6"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)

def build_docx(output_path):
    doc = Document()
    
    # Page setup: Standard 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Methods: Murine Spaceflight Transcriptomic Atlas & DNA Damage Repair Profiling")
        r_hdr.font.name = "Arial"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        
        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ftr = p_ftr.add_run("Methods Section — NASA OSDR Multi-Tissue Spaceflight Dataset (Teacher-Aligned Release)")
        r_ftr.font.name = "Arial"
        r_ftr.font.size = Pt(8.5)
        r_ftr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        
    COLOR_PRIMARY = RGBColor(0x1F, 0x4E, 0x78)
    COLOR_SECONDARY = RGBColor(0x2E, 0x75, 0xB6)
    COLOR_BODY = RGBColor(0x26, 0x26, 0x26)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Methods")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("Multi-Tissue Murine Spaceflight Transcriptomic Atlas and DNA Damage Repair Pathway Profiling")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_SECONDARY
    
    def add_sec_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        return p

    def add_body_p(content, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.22
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4.5)
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_BODY
        if isinstance(content, list):
            for item in content:
                if isinstance(item, str) and item.startswith("<m:oMath"):
                    m_el = parse_xml(item)
                    p._p.append(m_el)
                else:
                    r = p.add_run(item)
                    r.font.name = "Calibri"
                    r.font.size = Pt(10)
                    r.font.color.rgb = COLOR_BODY
        else:
            r = p.add_run(content)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.color.rgb = COLOR_BODY
        return p

    def add_bullet_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.18
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.left_indent = Inches(0.25)
        r_bullet = p.add_run("• ")
        r_bullet.font.name = "Calibri"
        r_bullet.font.size = Pt(10)
        r_bullet.font.color.rgb = COLOR_PRIMARY
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_BODY
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_BODY
        return p

    # ---------------------------------------------------------------------------
    # Section 1: Metadata & Experimental Design (Supplementary Table S1)
    # ---------------------------------------------------------------------------
    add_sec_heading("1. Experimental Design, Animal Cohort Compilation, and Multidimensional Metadata")
    
    add_body_p(
        "To establish an organism-wide, multi-tissue transcriptomic atlas of spaceflight adaptations and DNA damage repair (DDR) "
        "alterations under environmental stressors, transcriptomic profiles and experimental metadata were curated from the NASA Open Science "
        "Data Repository (OSDR, https://osdr.nasa.gov/). The analytical cohort encompasses 48 transcriptomic accessions (OSD datasets) "
        "across 59 primary analysis units, comprising 761 accession-scoped biospecimen records spanning 26 anatomical tissues (Supplementary Table S1). "
        "Ensembl Gene IDs serve as stable, unique computational identifiers throughout all analytical workflows, with gene symbols retained "
        "for biological annotation and cross-platform reference.",
        bold_prefix="Cohort compilation and biospecimen curation. "
    )
    
    add_body_p(
        "Biological and study metadata were programmatically curated from the NASA OSDR RESTful API and ISA-Tab archive files. "
        "The finalized analytical dataset encompasses 360 spaceflight biospecimens and 401 ground control records "
        "(comprising flight-matched habitat ground controls and vivarium controls). To capture mission-specific environmental dynamics "
        "across spaceflight missions, the cohort comprises "
        "13 distinct mission-level comparisons across 12 mission clusters (retaining SpaceX-8 and SpaceX-9 comparisons within OSD-162 as separate "
        "analytical units). Crucial biological attributes, including anatomical tissue (26 distinct tissues), biological sex (paired male and "
        "female cohorts), chronological age (15 discrete age brackets), and technical replicate designations, were harmonized into an audited reference "
        "matrix with 54 structured attributes to ensure complete sample-to-study traceability (Supplementary Table S1).",
        bold_prefix="Multidimensional metadata standardization and cohort governance. "
    )
    
    add_body_p(
        "Animal cohort allocation and multi-parameter stratification across the 26 anatomical tissues are systematically organized across mission clusters, "
        "sex-stratified female and male cohorts, and discrete chronological age brackets (Figure 1b, Supplementary Table S1). Anatomical tissues are ordered in accordance "
        "with cross-tissue DDR profiling panels, detailing accession-scoped unique donor counts and biospecimen distributions across spaceflight and ground "
        "control conditions to ensure robust cohort representation without confounding.",
        bold_prefix="Stratified experimental atlas organization. "
    )

    # ---------------------------------------------------------------------------
    # Section 2: Functional Pathway Enrichment & GO (Supplementary Tables S5, S6)
    # ---------------------------------------------------------------------------
    add_sec_heading("2. Functional Pathway Enrichment and Cross-Tissue Gene Ontology Profiling")
    
    add_body_p(
        [
            "To evaluate systemic pathway-level perturbations across anatomical systems under spaceflight environmental stressors, cross-tissue "
            "functional pathway enrichment was conducted using pre-ranked gene set enrichment analysis via fgseaMultilevel (fgsea package in R 4.4.3; "
            "random seed = 25). For each accession–tissue comparison, genes were ranked based on flight-versus-ground differential-expression statistics "
            "curated from OSDR. When available, repository test statistics were used directly; otherwise, signed standard normal deviates were computed from "
            "nominal two-sided P values and the sign of log2 fold change (log2FC): ",
            latex_to_omath(r"Z = -\Phi^{-1}\left(\frac{P}{2}\right) \times \operatorname{sign}(\log_2 \mathrm{FC})"),
            ". For duplicate Ensembl Gene IDs within a comparison, statistics were collapsed by median values, and ties were resolved deterministically by "
            "Ensembl ID. Gene set enrichment parameters were set to minSize = 5 for Gene Ontology terms and minSize = 10 for Hallmark terms, maxSize = 5,000, eps = ",
            latex_to_omath(r"10^{-10}"),
            ", with the standard score type."
        ],
        bold_prefix="Pre-ranked gene set enrichment analysis. "
    )
    
    add_body_p(
        "Enrichment profiling focused on 15 predefined Gene Ontology (GO) terms spanning 12 Biological Process and 3 Cellular Component terms "
        "relevant to genomic integrity and cellular stress: response to DNA damage stimulus (GO:0006974), DNA repair (GO:0006281), signal transduction "
        "in response to DNA damage (GO:0042770), intrinsic apoptotic signaling pathway (GO:0008630), DNA damage tolerance (GO:0006301), "
        "telomere maintenance in response to DNA damage (GO:0043247), chromosome, telomeric region (GO:0000781), cell cycle (GO:0007049), "
        "DNA replication (GO:0006260), DNA-templated transcription (GO:0006351), translation (GO:0006412), mitochondrion (GO:0005739), "
        "sensory perception of mechanical stimulus (GO:0050954), cytoskeleton (GO:0005856), and cell death (GO:0008219). Term memberships were "
        "constructed from the MGI MOD-GAF and go-basic ontology using relation-closure mapped to Ensembl release-116. In parallel, global transcriptional "
        "shifts were evaluated using the native mouse MSigDB Hallmark collection (MSigDB 2026.1.Mm; Figure S1, Supplementary Table S6).",
        bold_prefix="Curated Gene Ontology terms and Hallmark collections. "
    )
    
    add_body_p(
        "To avoid dominance by missions with redundant sub-studies, unit-level normalized enrichment scores (NES) within each tissue were first "
        "aggregated within mission clusters and then averaged with equal mission weights. Statistical significance was evaluated using exact two-sided "
        "sign-flip permutation tests across mission-level NES values, complemented by signed Stouffer meta-Z P values. False discovery rates were controlled "
        "by applying the Benjamini–Hochberg procedure separately across the 390 tissue–term cells (26 tissues × 15 GO terms) for each P-value family. "
        "Figure 1c presents the complete 26 × 15 mission-equal NES matrix (Supplementary Table S5). Supplementary Figure S1 displays the top 12 Hallmark "
        "pathways exhibiting the highest global positive and negative enrichment across missions with complete cohort coverage (Supplementary Table S6).",
        bold_prefix="Mission-equal NES and statistical meta-testing. "
    )

    # ---------------------------------------------------------------------------
    # Section 3: Sample-Level Association & Fisher-z (Supplementary Tables S7, S8)
    # ---------------------------------------------------------------------------
    add_sec_heading("3. Sample-Level Pathway Association and Cross-Tissue Fisher-z Integration")
    
    add_body_p(
        [
            "To quantify coordinated pathway co-regulation and evaluate whether DDR and metabolic pathways exhibit shared transcriptional trajectories "
            "across individual spaceflight animals, a ground-referenced pathway co-response framework was constructed. For each accession–tissue unit, "
            "flight specimen k was standardized against the baseline mean of matched ground control specimens: ",
            latex_to_omath(r"\mathrm{Sample\_log2FC}(g, k) = \log_2(x_{\mathrm{flight}}(g, k)) - \operatorname{mean}_{j}[\log_2(x_{\mathrm{ground}}(g, j))]"),
            ", where x represents normalized source expression values. Sample-level pathway activity scores were computed as "
            "the arithmetic mean of log2FC values across all detected member Ensembl IDs for each of the 15 predefined GO terms."
        ],
        bold_prefix="Ground-referenced sample pathway scoring. "
    )
    
    add_body_p(
        [
            "Within each of the 26 anatomical tissues, pairwise Spearman rank correlation coefficients (",
            latex_to_omath(r"\rho"),
            ") were calculated among the 15 pathway scores across individual flight specimens (Supplementary Table S8). To synthesize a consensus cross-tissue co-response architecture, correlation "
            "coefficients were transformed using the Fisher-z transformation: ",
            latex_to_omath(r"z = \operatorname{arctanh}(\rho)"),
            ", with bounded inputs within ",
            latex_to_omath(r"[-1 + 10^{-12},\, 1 - 10^{-12}]"),
            " and ",
            latex_to_omath(r"|z| \le 5.0"),
            ". Cross-tissue integration was performed by computing the equal-weight mean ",
            latex_to_omath(r"\bar{z}"),
            " across the 23 tissues with sufficient flight sample sizes (",
            latex_to_omath(r"n_{\mathrm{flight}} \ge 5"),
            "). Three tissues with limited sample numbers—Colon (n = 4), Heart (n = 3), and Femoral skin (n = 2)—were evaluated in "
            "tissue-specific correlation profiles (Supplementary Table S8) and excluded from the cross-tissue pooled integration. The consensus correlation "
            "matrix was obtained by back-transformation: ",
            latex_to_omath(r"\bar{\rho} = \tanh(\bar{z})"),
            ", yielding the consensus 15 × 15 cross-tissue pathway co-response matrix depicted in Figure 2 (Supplementary Table S7)."
        ],
        bold_prefix="Intra-tissue correlation and Fisher-z meta-integration. "
    )

    # ---------------------------------------------------------------------------
    # Section 4: DDR Repertoires & Common-Direction (Supplementary Table S2)
    # ---------------------------------------------------------------------------
    add_sec_heading("4. Organ-Specific DDR Member-Gene Repertoires and Common-Direction Prioritization")
    
    add_body_p(
        [
            "To determine whether shared or tissue-specific gene repertoires drive spaceflight DDR perturbations, binary detection matrices of Ensembl IDs "
            "belonging to the core DDR term (GO:0006974) were evaluated. A member gene was defined as present when its normalized source expression was finite "
            "and greater than zero. Combinatorial overlap and intersection sets of expressed DDR genes were quantified across independent flight comparisons "
            "within Thymus (4 comparisons; Figure 3a, Supplementary Table S2), Kidney (5 accessions; Figure S2a, Supplementary Table S2), as well as across tissue cohorts "
            "exhibiting pronounced positive enrichment (",
            latex_to_omath(r"\mathrm{NES} > 1"),
            "; Figure S3a, Supplementary Table S2) and negative enrichment (",
            latex_to_omath(r"\mathrm{NES} < -1"),
            "; Figure S4a, Supplementary Table S2)."
        ],
        bold_prefix="UpSet intersection topology of DDR member genes. "
    )
    
    add_body_p(
        [
            "To identify core responsive genes exhibiting reproducible, unidirectional regulation under spaceflight stress across parallel studies, candidate "
            "genes were prioritized based on directional concordance. Mission-level differential expression statistics were corrected using Benjamini–Hochberg "
            "FDR (threshold FDR ≤ 0.05 for Thymus and Kidney; FDR ≤ 0.10 for the high- and low-NES tissue cohorts). Qualifying genes were required to display "
            "identical non-zero log2FC signs across all parallel comparisons within the cohort. Concordant genes were then ranked by their minimum effect "
            "magnitude across comparisons: ",
            latex_to_omath(r"\mathrm{Score}_g = \min_{c} |\log_2 \mathrm{FC}(g, c)|"),
            ". This prioritization identified 30 top concordant genes in Thymus (Figure 3b, Supplementary Table S2), 12 in Kidney (Figure S2b, Supplementary Table S2), 5 in high-NES tissues (Figure S3b, Supplementary Table S2), and 11 in low-NES tissues (Figure S4b, Supplementary Table S2), compiling comparison-specific log2FC values, FDR significance, and concordance scores."
        ],
        bold_prefix="Common-direction concordance and candidate prioritization. "
    )

    # ---------------------------------------------------------------------------
    # Section 5: DSB Repair Architecture (Supplementary Tables S4, S9)
    # ---------------------------------------------------------------------------
    add_sec_heading("5. Double-Strand Break Repair Architecture and Spaceflight Meta-Analysis")
    
    add_body_p(
        [
            "To characterize baseline expression patterns and tissue divergence in double-strand break (DSB) repair machinery, flight transcriptomic profiles "
            "were normalized using YARN tissue-aware quantile smoothing (yarn::normalizeTissueAware; groups = tissue, normalizationMethod = 'qsmooth', "
            "window = 0.05, log = TRUE). Normalization was conducted across a background universe of 12,462 Ensembl genes, 360 flight biospecimens, "
            "59 analysis units, and 26 tissues. For DSB architecture analysis, 29 curated pathway components were compiled across non-homologous end joining "
            "(NHEJ, 7 genes), homologous recombination (HR, 14 genes), HR/A-EJ shared components (3 genes), and alternative end joining (A-EJ, 5 genes). "
            "Expression profiles were sequentially aggregated by biospecimen, analysis unit, mission, and tissue. Pairwise tissue dissimilarity was calculated "
            "as ",
            latex_to_omath(r"D(T_a, T_b) = 1 - \rho_{\mathrm{Spearman}}"),
            " on tissue-mean YARN normalized log2 profiles, followed by average-linkage hierarchical clustering to reconstruct anatomical relationships "
            "and evaluate cross-tissue co-expression clustering (Figure 4a)."
        ],
        bold_prefix="Tissue-aware normalization and pathway clustering. "
    )
    
    add_body_p(
        [
            "To quantify physiological expression distributions and tissue-specific abundance of key DSB repair machinery, representative marker genes were "
            "quantified on normalized log2 scale (as well as back-transformed linear scale: ",
            latex_to_omath(r"\mathrm{Expression}_{\mathrm{linear}} = 2^{\mathrm{YARNNormalizedLog2}} - 1"),
            ") across 360 flight biospecimens in 26 anatomical tissues (Supplementary Table S4):"
        ],
        bold_prefix="Baseline expression profiling of DSB machinery. "
    )
    
    add_bullet_p(
        "Nhej1 / XLF (ENSMUSG00000026162), Paxx (ENSMUSG00000047617), Xrcc6 / Ku70 (ENSMUSG00000022471), and Xrcc5 / Ku80 (ENSMUSG00000026187).",
        bold_prefix="NHEJ (Supplementary Table S4, 4 genes): "
    )
    add_bullet_p(
        "Brca1 (ENSMUSG00000017146), Bard1 (ENSMUSG00000026196), Blm (ENSMUSG00000030528), and Rad51 (ENSMUSG00000027323).",
        bold_prefix="HR (Supplementary Table S4, 4 genes): "
    )
    add_bullet_p(
        "Parp1 (ENSMUSG00000026496), Polq (ENSMUSG00000034206), Lig1 (ENSMUSG00000056394), and Lig3 (ENSMUSG00000020697).",
        bold_prefix="A-EJ (Supplementary Table S4, 4 genes): "
    )
    
    add_body_p(
        "In the stratified boxplots (Figure 4b–d), the central solid line denotes the sample mean (μ), the lower and upper box hinges represent the first "
        "(Q1, 25%) and third (Q3, 75%) quartiles, and whiskers extend to the furthest observations within 1.5 × IQR. Individual flight biospecimens "
        "(360 samples) are overlaid as jittered points (dodge width = 0.80, jitter width = 0.11, seed = 25), with dashed lines connecting tissue means. "
        "Between-tissue expression heterogeneity for each repair gene across the 360 flight biospecimens was formally evaluated using ANOVA across "
        "anatomical tissues, with summary statistics, F-test values, and P values compiled in Supplementary Table S4.",
        bold_prefix="Boxplot geometry and tissue variance. "
    )
    
    add_body_p(
        [
            "To assess spaceflight transcriptional shifts across the complete 29-component DSB repertoire, multi-mission meta-analysis was performed across "
            "all 26 tissues (Figure 6a, Supplementary Table S9). The overall spaceflight effect size was calculated as the median of mission-level median log2FC values. "
            "Meta-analytic significance across missions was determined using signed Stouffer's Z-trend tests. Significant perturbations with nominal tissue-meta "
            "P < 0.05 are presented as meta-dotplots, wherein dot color reflects the direction of regulation (red, spaceflight-upregulated; blue, "
            "spaceflight-downregulated) and dot area scales with effect magnitude (",
            latex_to_omath(r"|\log_2 \mathrm{FC}|"),
            ") (Figure 6a, Supplementary Table S9)."
        ],
        bold_prefix="Spaceflight-induced meta-differential expression analysis. "
    )

    # ---------------------------------------------------------------------------
    # Section 6: SSB Repair Architecture (Supplementary Tables S4, S10)
    # ---------------------------------------------------------------------------
    add_sec_heading("6. Single-Strand Break Repair Architecture and Spaceflight Meta-Analysis")
    
    add_body_p(
        [
            "To establish baseline expression architectures across single-strand break (SSB) repair mechanisms, a curated panel of 45 components was compiled "
            "across the four major SSB pathways: base excision repair (BER, 6 genes), nucleotide excision repair (NER, 18 genes), mismatch repair (MMR, 5 genes), "
            "and the Fanconi anemia pathway (FA, 16 genes). Biospecimen expression profiles were normalized using YARN tissue-aware quantile smoothing (qsmooth), "
            "followed by successive aggregation to tissue means. Pairwise tissue distances were computed as ",
            latex_to_omath(r"D(T_a, T_b) = 1 - \rho_{\mathrm{Spearman}}"),
            " and resolved using average-linkage hierarchical clustering to reveal cross-tissue organizational dendrograms (Figure 5a)."
        ],
        bold_prefix="Pathway definition and tissue-aware clustering. "
    )
    
    add_body_p(
        [
            "Tissue-specific expression distributions of key SSB components were evaluated on normalized log2 scale (as well as back-transformed linear scale: ",
            latex_to_omath(r"\mathrm{Expression}_{\mathrm{linear}} = 2^{\mathrm{YARNNormalizedLog2}} - 1"),
            ") across 360 flight biospecimens in 26 tissues (Supplementary Table S4):"
        ],
        bold_prefix="Baseline expression profiling of SSB machinery. "
    )
    
    add_bullet_p(
        "Ung (ENSMUSG00000029591), Ogg1 (ENSMUSG00000030271), and Neil1 (ENSMUSG00000032298).",
        bold_prefix="BER (Supplementary Table S4, 3 genes): "
    )
    add_bullet_p(
        "Xpc (ENSMUSG00000030094), Rad23b (ENSMUSG00000028426), Cetn2 (ENSMUSG00000031347), Ddb1 (ENSMUSG00000024740), Ddb2 (ENSMUSG00000002109), "
        "Ercc6 (ENSMUSG00000054051), and Ercc8 (ENSMUSG00000021694), representing core global genome repair (XPC–RAD23B–CETN2, DDB1–DDB2) "
        "and transcription-coupled repair (ERCC6, ERCC8) subcomplexes.",
        bold_prefix="NER (Supplementary Table S4, 7 genes): "
    )
    add_bullet_p(
        "Msh2 (ENSMUSG00000024151) and Msh3 (ENSMUSG00000014850).",
        bold_prefix="MMR (Supplementary Table S4, 2 genes): "
    )
    add_bullet_p(
        "Fancd2 (ENSMUSG00000034023) and Fanci (ENSMUSG00000039187).",
        bold_prefix="FA (Supplementary Table S4, 2 genes): "
    )
    
    add_body_p(
        "Boxplot geometry, sample jittering (seed = 25), and tissue-mean connectors followed the DSB specifications (Figure 5b–e). "
        "Between-tissue expression heterogeneity for each SSB repair gene across the 360 flight biospecimens was formally evaluated "
        "using ANOVA across anatomical tissues, with detailed statistics compiled in Supplementary Table S4.",
        bold_prefix="Boxplot geometry and tissue variance. "
    )
    
    add_body_p(
        [
            "To determine transcriptional perturbations across the 45 SSB components in spaceflight, multi-mission meta-analysis was conducted across "
            "all 26 tissues (Figure 6b, Supplementary Table S10). Meta-effect sizes were calculated as the median of mission-level median log2FC values, and statistical "
            "significance was evaluated using signed Stouffer's Z-trend tests. Associations meeting nominal tissue-meta P < 0.05 are displayed as meta-dotplots "
            "(Figure 6b, Supplementary Table S10), with dot color indicating regulation direction and dot size encoding effect magnitude (",
            latex_to_omath(r"|\log_2 \mathrm{FC}|"),
            ")."
        ],
        bold_prefix="Spaceflight-induced meta-differential expression analysis. "
    )

    # ---------------------------------------------------------------------------
    # Section 7: Detection Audit (Supplementary Table S3)
    # ---------------------------------------------------------------------------
    add_sec_heading("7. Multi-Pathway Member-Gene Detection Audit Across Tissues")
    
    add_body_p(
        "To confirm the detection coverage and analytical integrity of the DNA damage repair machinery across the entire spaceflight dataset, expression "
        "detection was audited for 257 non-redundant Ensembl Gene IDs comprising 335 pathway memberships across the seven curated repair modules: "
        "BER (54 memberships), NER (53), MMR (27), FA (39), HR (135), A-EJ (9), and NHEJ (18). A gene was classified as detected in a given biospecimen "
        "when its normalized source expression value was finite and greater than zero. Detection counts were calculated for each of the 761 biospecimens "
        "across the 26 anatomical tissues, stratified by flight and ground conditions, to verify robust transcriptomic coverage across experimental "
        "groups (Figure S5, Supplementary Table S3).",
        bold_prefix="Pathway coverage and biospecimen detection integrity. "
    )

    # ---------------------------------------------------------------------------
    # Section 8: Environment and Reproducibility
    # ---------------------------------------------------------------------------
    add_sec_heading("8. Statistical Analysis, Software Environment, and Computational Reproducibility")
    
    add_body_p(
        "All bioinformatics data processing, statistical meta-analyses, and visualization workflows were developed in R 4.4.3 and Python "
        "(v3.13/v3.14). Key pinned R packages include data.table (v1.18.2.1), ggplot2 (v4.0.2), patchwork (v1.3.2), scales (v1.4.0), and "
        "ggVennDiagram (v1.5.7). Python workflows utilized matplotlib (v3.10.8), numpy (v2.5.0), and pandas (v3.0.2). Multiple-testing adjustments "
        "throughout the study were performed using the Benjamini–Hochberg procedure to control the false discovery rate. Between-tissue expression "
        "heterogeneity for linear and log2 distributions was assessed via ANOVA across anatomical tissues.",
        bold_prefix="Statistical standards and computational environment. "
    )
    
    add_body_p(
        "To ensure complete computational reproducibility, all stochastic algorithms, permutation routines, and point jittering were locked to "
        "random seed 25. All analytical code, statistical models, and visualization pipelines are version-controlled under Git and deposited in the "
        "study repository (https://github.com/gzy2520/mouse-spaceflight-transcriptome-publication). Released analytical tables, numerical datasets, and input "
        "manifests are verified by cryptographic SHA-256 checksums (provenance/reference_sha256.csv), enabling end-to-end auditability from raw OSDR inputs to "
        "final analytical tables and publication outputs.",
        bold_prefix="Reproducibility protocols, seed locking, and version governance. "
    )

    # ---------------------------------------------------------------------------
    # Section 9: Reference Matrix (Supplementary Table Cross-Reference)
    # ---------------------------------------------------------------------------
    add_sec_heading("9. Analytical Workflow and Supplementary Table Cross-Reference Matrix")
    
    add_body_p(
        "The following reference matrix systematically connects each Supplementary Table and primary analytical module to its analytical scope, "
        "input dataset, mathematical model, and reproducible source script or data contract in sequential order, strictly harmonized with the "
        "manuscript figure citations:",
        bold_prefix="Provenance and workflow cross-reference. "
    )
    
    table_data = [
        ["Supplementary Table", "Analytical Scope", "Input Cohort / Sample Size", "Statistical Model / Algorithm", "Source Script / Reference Contract"],
        ["Supplementary Table S1", "Animal Cohort Compilation and Multidimensional Metadata Atlas", "761 samples, 48 accessions, 49 cohorts, 26 tissues (360 flight, 401 ground)", "Harmonized 54-attribute clinical and mission metadata, categorical age/sex stratification", "provenance/Fig_1b_mouse_metadata_grouped.csv; Fig_6_mouse_sample_metadata_complete_audit.csv; R/final_figures/02_fig1b_metadata_bubble.R"],
        ["Supplementary Table S2", "Spaceflight-Induced Co-Directionally Regulated DDR Genes and Repertoire Overlap", "Thymus (4 cohorts), Kidney (5 accessions), Up-cohort (9 tissues), Down-cohort (6 tissues)", "GO:0006974 member Ensembl-ID overlap (>94% core); BH-FDR and concordant-direction ranking (Thymus 30 DEGs, Kidney 12 DEGs, Up 5 DEGs, Down 11 DEGs)", "data/publication_input/common_direction/ (A, B, C, D); data/publication_input/venn/; results/tables/04_membership_matrix_..."],
        ["Supplementary Table S3", "Seven DNA Damage Repair Pathways Gene Repertoire Sizes and Multi-Tissue Detection Audit", "761 biospecimens across 26 tissues × 7 repair pathways (BER, NER, MMR, FA, HR, A-EJ, NHEJ)", "Curated repertoire sizes (HR 135, BER 54, NER 53, FA 39, MMR 27, NHEJ 18, A-EJ 9); finite normalized source-expression detection audits", "release/tables/08_Table_S3_pathway_repertoire_sizes.csv; 09_Table_S3_per_tissue_sample_summary.csv; 01_seven_pathway_member_gene_counts_per_source_sample_wide.csv"],
        ["Supplementary Table S4", "Baseline Expression Profiles and ANOVA Heterogeneity of Core Repair Genes Across Tissues", "360 Flight samples across 26 tissues × 26 core repair genes (NHEJ, HR, A-EJ, BER, NER, MMR, FA)", "Normalized log2 (and linear) expression distributions, tissue means, and one-way/two-way ANOVA across anatomical tissues", "release/tables/log/02_gene_one_way_anova_log_summary.csv; 03_tissue_gene_mean_log_qsmooth_summary.csv; R/final_figures/09_pathway_gene_tissue_boxplots_log.R"],
        ["Supplementary Table S5", "Cross-Tissue Gene Ontology Enrichment Profiling Matrix", "26 tissues × 15 GO terms (12 BP, 3 CC)", "Mission-equal mean NES; exact sign-flip permutation and secondary Stouffer meta-Z statistics, BH FDR", "data/publication_input/go/04_tissue_statistics_concrete_terms_and_context.csv; R/final_figures/03_fig1c_fig2_go.R"],
        ["Supplementary Table S6", "MSigDB Hallmark Pathway GSEA Overview Across Missions", "13 missions with complete cohort coverage", "Pre-ranked fgseaMultilevel; top 12 global Hallmark pathways (6 positive, 6 negative NES)", "data/publication_input/hallmark/01_global_top_term_figure_selection.csv; R/figures/02_hallmark_gsea.R"],
        ["Supplementary Table S7", "Cross-Tissue Consensus Pathway Association Matrix", "23 eligible tissues (n_flight ≥ 5; 351 flight samples)", "Ground-referenced sample log2FC, intra-tissue Spearman rank correlation, equal-weight Fisher-z pooling", "data/publication_input/go/log2fc_spearman_20260902/04_overall_fisher_z_spearman_rho_matrix.csv; python/plot_log2fc_per_tissue_pathway_spearman_20260902.py"],
        ["Supplementary Table S8", "Per-Tissue Sample-Level Pathway Spearman Correlation Profiles", "26 individual tissue matrices across flight samples", "Pairwise 15 × 15 Spearman rank correlation matrices across flight specimens per tissue", "data/publication_input/go/log2fc_spearman_20260902/03_tissue_pairwise_spearman_log2fc_long.csv; python/build_log2fc_per_tissue_pathway_spearman_20260902.py"],
        ["Supplementary Table S9", "Double-Strand Break (DSB) Repair Multi-Mission Spaceflight Meta-Analysis Matrix", "26 tissues × 29 DSB repair components", "Median of mission medians log2FC; signed Stouffer meta-Z P values; nominal P < 0.05 indicators", "data/publication_input/meta/08_essential_components_counts_log2FC_tissue_matrix.csv; R/final_figures/06_meta_dotplots.R"],
        ["Supplementary Table S10", "Single-Strand Break (SSB) Repair Multi-Mission Spaceflight Meta-Analysis Matrix", "26 tissues × 45 SSB repair components", "Median of mission medians log2FC; signed Stouffer meta-Z P values; nominal P < 0.05 indicators", "release/tables/01_SSB_tissue_meta_log2FC_pvalue_matrix_without_Neil2.csv; R/final_figures/06_meta_dotplots.R"]
    ]
    
    table = doc.add_table(rows=len(table_data), cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(1.2), Inches(1.3), Inches(1.3), Inches(1.5), Inches(1.2)]
    for i, col in enumerate(table.columns):
        col.width = col_widths[i]
        
    set_three_line_borders(table)
    
    for row_idx, row_content in enumerate(table_data):
        row = table.rows[row_idx]
        is_header = (row_idx == 0)
        
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for col_idx, text in enumerate(row_content):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            if is_header:
                set_cell_background(cell, "EBF1F5")
                set_cell_margins(cell, top=100, bottom=100, left=70, right=70)
            else:
                if row_idx % 2 == 1:
                    set_cell_background(cell, "FAFCFD")
                set_cell_margins(cell, top=55, bottom=55, left=70, right=70)
                
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            run = p.add_run(text)
            run.font.name = "Calibri"
            if is_header:
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = COLOR_PRIMARY
            else:
                run.font.size = Pt(8.0)
                run.font.color.rgb = COLOR_BODY

    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:val"), "bestFit")
        zoom.set(qn("w:percent"), "100")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Successfully generated Supplementary Table-Aligned Methods document: {output_path}")

if __name__ == "__main__":
    repo_root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    out_repo = os.path.join(repo_root, "manuscript_methods_mouse_spaceflight_transcriptome.docx")
    build_docx(out_repo)
    
    out_desktop_revised = "/Users/gzy2520/Desktop/mouse_Methods_revised.docx"
    build_docx(out_desktop_revised)

    out_desktop = "/Users/gzy2520/Desktop/mouse_Methods.docx"
    build_docx(out_desktop)

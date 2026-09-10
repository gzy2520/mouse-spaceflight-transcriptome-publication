#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build complete publication delivery tables aligned with Teacher's Manuscript & Methods:
- S1 to S4: Teacher's designated core tables (Metadata, Overlap/DEGs, 7 Pathways Audit, Core Genes ANOVA).
- S5 to S8: Analytical results tables for remaining manuscript sections (GO/Hallmark GSEA, Fisher-z Network, DSB Meta, SSB Meta).
- Nomenclature: Strictly 'Table S' notation, removing 'Supplementary'.
- Formatting: Streamlined delivery (no internal parameters/audit flags), completely unstyled (no fill colors).
- Output: /Users/gzy2520/Desktop/Supplementary_Tables/
"""

import os
import gzip
import shutil
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border
from openpyxl.utils import get_column_letter

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DESKTOP_DIR = "/Users/gzy2520/Desktop/Supplementary_Tables"
CSV_DIR = os.path.join(DESKTOP_DIR, "csv")

FONT_FAMILY = "Calibri"
HEADER_FONT = Font(name=FONT_FAMILY, size=11, bold=True, color="000000")
DATA_FONT = Font(name=FONT_FAMILY, size=10, color="000000")

def format_worksheet_unstyled(ws, df, is_first_col_bold=False):
    """
    Apply clean, unstyled publication formatting:
    - NO background fills (completely plain/unstyled)
    - Default native gridlines enabled
    - Header in 11pt bold black
    - Numeric values aligned right with standard formats
    - Auto-adjusted column widths
    """
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"
    
    ws.row_dimensions[1].height = 24
    for col_num in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = HEADER_FONT
        cell.fill = PatternFill(fill_type=None)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
        cell.border = Border()

    num_rows = len(df)
    num_cols = len(df.columns)
    for row_idx in range(num_rows):
        row_num = row_idx + 2
        ws.row_dimensions[row_num].height = 19
        for col_idx in range(num_cols):
            cell = ws.cell(row=row_num, column=col_idx + 1)
            cell.font = DATA_FONT
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()
            
            val = cell.value
            if isinstance(val, (int,)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0"
            elif isinstance(val, (float,)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if abs(val) < 0.001 and val != 0:
                    cell.number_format = "0.00E+00"
                else:
                    cell.number_format = "0.0000"
            else:
                if col_idx == 0 and is_first_col_bold:
                    cell.font = Font(name=FONT_FAMILY, size=10, bold=True, color="000000")
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            v_str = str(cell.value or '')
            if len(v_str) > max_len:
                max_len = len(v_str)
        col_width = max(max_len + 3, 11)
        if col_width > 50:
            col_width = 50
        ws.column_dimensions[col_letter].width = col_width

def write_excel_sheets(output_xlsx_path, sheets_dict):
    """Write multiple dataframes to an unstyled clean Excel workbook."""
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)
    
    for sheet_name, df in sheets_dict.items():
        ws = wb.create_sheet(title=sheet_name[:31])
        for col_idx, col_name in enumerate(df.columns, start=1):
            ws.cell(row=1, column=col_idx, value=col_name)
            
        for row_idx, row_data in enumerate(df.itertuples(index=False), start=2):
            for col_idx, value in enumerate(row_data, start=1):
                if pd.isna(value):
                    ws.cell(row=row_idx, column=col_idx, value="")
                else:
                    ws.cell(row=row_idx, column=col_idx, value=value)
                    
        format_worksheet_unstyled(ws, df)
        
    wb.save(output_xlsx_path)
    print(f"Generated: {output_xlsx_path}")

def save_csvs(table_prefix, sheets_dict):
    """Save each sheet as a standalone CSV file."""
    for sheet_name, df in sheets_dict.items():
        clean_name = sheet_name.replace(" ", "_").replace("/", "_")
        csv_path = os.path.join(CSV_DIR, f"{table_prefix}_{clean_name}.csv")
        df.to_csv(csv_path, index=False)
        print(f"  CSV: {csv_path}")

def clean_deg_df(df):
    core_cols = ['ensembl_id', 'display_symbol', 'in_DNA_damage_response_GO0006974', 'in_DNA_repair_GO0006281', 'common_direction']
    comp_cols = [c for c in df.columns if any(c.startswith(prefix) for prefix in ['log2fc__', 'wald_z__', 'sample_fdr__'])]
    stat_cols = ['mean_log2fc', 'min_abs_log2fc', 'final_integrated_z', 'final_integrated_p', 'final_fdr', 'pooled_rank']
    selected = [c for c in core_cols + comp_cols + stat_cols if c in df.columns]
    res = df[selected].rename(columns={'display_symbol': 'gene_symbol'})
    return res

def clean_upset_df(df):
    set_cols = [c for c in df.columns if c.startswith('set_')]
    res = df[['ensembl_id', 'display_symbol'] + set_cols].rename(columns={'display_symbol': 'gene_symbol'})
    res['intersection_degree'] = res[set_cols].sum(axis=1)
    return res

def main():
    if os.path.exists(DESKTOP_DIR):
        shutil.rmtree(DESKTOP_DIR)
    os.makedirs(CSV_DIR, exist_ok=True)

    print("=== Step 1: Generating Table S1 (Animal Cohort Compilation & Sample Metadata) ===")
    df_s1_cohorts_raw = pd.read_csv(os.path.join(REPO_ROOT, "reproduced_final_result_20260908_tree_axis_order_complete/provenance/Fig_1b_mouse_metadata_grouped.csv"))
    df_s1_cohorts = df_s1_cohorts_raw[['mission_cluster', 'Group', 'sex', 'age_label', 'n_samples', 'n_mice', 'n_flight', 'n_ground', 'n_accessions', 'accessions', 'n_analysis_units', 'analysis_units']].rename(columns={
        'mission_cluster': 'Mission_Cluster',
        'Group': 'Tissue',
        'sex': 'Sex',
        'age_label': 'Age_Group',
        'n_samples': 'Total_Samples',
        'n_mice': 'Unique_Mice',
        'n_flight': 'Flight_Samples',
        'n_ground': 'Ground_Control_Samples',
        'n_accessions': 'OSD_Accession_Count',
        'accessions': 'OSD_Accessions',
        'n_analysis_units': 'Analysis_Units_Count',
        'analysis_units': 'Analysis_Units'
    })

    df_s1_samples_raw = pd.read_csv(os.path.join(REPO_ROOT, "reproduced_final_result_20260908_tree_axis_order_complete/provenance/Fig_6_mouse_sample_metadata_complete_audit.csv"))
    df_s1_samples = df_s1_samples_raw[['accession', 'analysis_unit_id', 'tissue', 'normalized_sample_name', 'accession_scoped_mouse_id', 'analysis_sample_status', 'sample_is_technical_replicate', 'sex', 'age_label', 'mission_cluster', 'source_material_type', 'osdr_study_url']].rename(columns={
        'accession': 'OSD_Accession',
        'analysis_unit_id': 'Analysis_Unit_ID',
        'tissue': 'Tissue',
        'normalized_sample_name': 'Sample_Name',
        'accession_scoped_mouse_id': 'Mouse_ID',
        'analysis_sample_status': 'Condition',
        'sample_is_technical_replicate': 'Is_Technical_Replicate',
        'sex': 'Sex',
        'age_label': 'Age_Group',
        'mission_cluster': 'Mission_Cluster',
        'source_material_type': 'Assay_Type',
        'osdr_study_url': 'OSDR_Repository_URL'
    })

    s1_sheets = {
        "Cohort_Summary_Fig1b": df_s1_cohorts,
        "Sample_Metadata_Audit_761": df_s1_samples
    }
    s1_path = os.path.join(DESKTOP_DIR, "Table_S1_Cohort_and_Sample_Metadata.xlsx")
    write_excel_sheets(s1_path, s1_sheets)
    save_csvs("Table_S1", s1_sheets)

    print("\n=== Step 2: Generating Table S2 (Spaceflight DDR Overlap & Common DEGs) ===")
    df_s2_thymus_deg = clean_deg_df(pd.read_csv(os.path.join(REPO_ROOT, "data/publication_input/common_direction/B_thymus_DDR_common_direction_top30_min_abs_log2fc_direction_ordered.csv")))
    df_s2_thymus_upset = clean_upset_df(pd.read_csv(os.path.join(REPO_ROOT, "data/publication_input/venn/04_membership_matrix_thymus_four_comparisons.csv")))
    df_s2_kidney_deg = clean_deg_df(pd.read_csv(os.path.join(REPO_ROOT, "data/publication_input/common_direction/A_kidney_DDR_common_direction_top30_min_abs_log2fc_direction_ordered.csv")))
    df_s2_kidney_upset = clean_upset_df(pd.read_csv(os.path.join(REPO_ROOT, "data/publication_input/venn/04_membership_matrix_kidney_five_datasets.csv")))
    df_s2_up_deg = clean_deg_df(pd.read_csv(os.path.join(REPO_ROOT, "data/publication_input/common_direction/D_NES_gt1_excl_Kidney_DDR_common_direction_top30_min_abs_log2fc_direction_ordered.csv")))
    df_s2_up_upset = clean_upset_df(pd.read_csv(os.path.join(REPO_ROOT, "results/tables/04_membership_matrix_up_tissues.csv")))
    df_s2_down_deg = clean_deg_df(pd.read_csv(os.path.join(REPO_ROOT, "data/publication_input/common_direction/C_NES_lt_minus1_excl_Lung_Thymus_DDR_common_direction_top30_min_abs_log2fc_direction_ordered.csv")))
    df_s2_down_upset = clean_upset_df(pd.read_csv(os.path.join(REPO_ROOT, "results/tables/04_membership_matrix_down_tissues.csv")))

    s2_sheets = {
        "Thymus_Top30_DEGs": df_s2_thymus_deg,
        "Thymus_DDR_Overlap_893": df_s2_thymus_upset,
        "Kidney_Common12_DEGs": df_s2_kidney_deg,
        "Kidney_DDR_Overlap_893": df_s2_kidney_upset,
        "Up_Cohort_Common5_DEGs": df_s2_up_deg,
        "Up_Cohort_DDR_Overlap_893": df_s2_up_upset,
        "Down_Cohort_Common11_DEGs": df_s2_down_deg,
        "Down_Cohort_DDR_Overlap_893": df_s2_down_upset
    }
    s2_path = os.path.join(DESKTOP_DIR, "Table_S2_DDR_Overlap_and_Common_DEGs.xlsx")
    write_excel_sheets(s2_path, s2_sheets)
    save_csvs("Table_S2", s2_sheets)

    print("\n=== Step 3: Generating Table S3 (Seven DNA Repair Pathways Repertoires & Detection Audit) ===")
    df_s3_sizes_raw = pd.read_csv(os.path.join(REPO_ROOT, "release/tables/08_Table_S3_pathway_repertoire_sizes.csv"))
    df_s3_sizes = df_s3_sizes_raw[['Pathway', 'MemberGeneCount']].rename(columns={'MemberGeneCount': 'Member_Gene_Count'})

    df_s3_tissue_raw = pd.read_csv(os.path.join(REPO_ROOT, "release/tables/09_Table_S3_per_tissue_sample_summary.csv"))
    df_s3_tissue = df_s3_tissue_raw[['Group', 'Pathway', 'n_unique_Ensembl_genes_in_GO_set', 'n_samples', 'n_flight_samples', 'n_ground_samples', 'min', 'q1', 'median', 'mean', 'q3', 'max']].rename(columns={
        'Group': 'Tissue',
        'n_unique_Ensembl_genes_in_GO_set': 'Pathway_Total_Genes',
        'n_samples': 'Total_Samples',
        'n_flight_samples': 'Flight_Samples',
        'n_ground_samples': 'Ground_Samples',
        'min': 'Min_Detected_Genes',
        'q1': 'Q1_Detected_Genes',
        'median': 'Median_Detected_Genes',
        'mean': 'Mean_Detected_Genes',
        'q3': 'Q3_Detected_Genes',
        'max': 'Max_Detected_Genes'
    })

    df_s3_samples_raw = pd.read_csv(os.path.join(REPO_ROOT, "release/tables/01_seven_pathway_member_gene_counts_per_source_sample_wide.csv"))
    df_s3_samples = df_s3_samples_raw[['Group', 'accession', 'analysis_unit_id', 'sample_status', 'sample_is_technical_replicate', 'n_BER', 'n_NER', 'n_MMR', 'n_FA', 'n_HR', 'n_AEJ', 'n_NHEJ']].rename(columns={
        'Group': 'Tissue',
        'accession': 'OSD_Accession',
        'analysis_unit_id': 'Analysis_Unit_ID',
        'sample_status': 'Condition',
        'sample_is_technical_replicate': 'Is_Technical_Replicate'
    })

    s3_sheets = {
        "Pathway_Repertoire_Sizes": df_s3_sizes,
        "Tissue_Detection_Summary": df_s3_tissue,
        "Sample_Pathway_Counts_761": df_s3_samples
    }
    s3_path = os.path.join(DESKTOP_DIR, "Table_S3_Seven_Pathways_Gene_Repertoires_and_Detection.xlsx")
    write_excel_sheets(s3_path, s3_sheets)
    save_csvs("Table_S3", s3_sheets)

    print("\n=== Step 4: Generating Table S4 (Core Repair Genes Baseline Expression & ANOVA Heterogeneity) ===")
    df_s4_gene_anova_raw = pd.read_csv(os.path.join(REPO_ROOT, "release/tables/log/02_gene_one_way_anova_log_summary.csv"))
    df_s4_gene_anova = df_s4_gene_anova_raw[['Pathway', 'EnsemblID', 'Symbol', 'Df_group', 'Df_residual', 'F_value', 'P_value', 'Significance']].rename(columns={'Symbol': 'Gene_Symbol'})

    df_s4_pathway_anova = pd.read_csv(os.path.join(REPO_ROOT, "release/tables/log/01_pathway_one_way_anova_log_summary.csv"))[['Pathway', 'Df_group', 'Df_residual', 'F_value', 'P_value', 'Significance']]

    df_s4_tissue_mean_raw = pd.read_csv(os.path.join(REPO_ROOT, "release/tables/log/03_tissue_gene_mean_log_qsmooth_summary.csv"))
    df_s4_tissue_mean = df_s4_tissue_mean_raw[['Pathway', 'Group', 'EnsemblID', 'Gene', 'Mean', 'SD', 'Q1', 'Q3', 'N']].rename(columns={
        'Group': 'Tissue',
        'Gene': 'Gene_Symbol',
        'Mean': 'Mean_Log2_Expression',
        'SD': 'SD_Log2_Expression',
        'Q1': 'Q1_Log2',
        'Q3': 'Q3_Log2',
        'N': 'Sample_Count'
    })

    with gzip.open(os.path.join(REPO_ROOT, "release/tables/log/04_sample_log_qsmooth_values_used.csv.gz"), "rt") as f:
        df_s4_samples_raw = pd.read_csv(f)
    df_s4_samples = df_s4_samples_raw[['Pathway', 'Group', 'EnsemblID', 'OfficialMouseSymbol', 'SampleID', 'accession', 'mission_cluster', 'YARNNormalizedLog2']].rename(columns={
        'Group': 'Tissue',
        'OfficialMouseSymbol': 'Gene_Symbol',
        'accession': 'OSD_Accession',
        'mission_cluster': 'Mission_Cluster',
        'YARNNormalizedLog2': 'Normalized_Log2_Expression'
    })

    s4_sheets = {
        "ANOVA_Core_Genes": df_s4_gene_anova,
        "ANOVA_Pathways": df_s4_pathway_anova,
        "Tissue_Gene_Mean_Log2": df_s4_tissue_mean,
        "Sample_Log2_Values": df_s4_samples
    }
    s4_path = os.path.join(DESKTOP_DIR, "Table_S4_Repair_Genes_Expression_and_ANOVA.xlsx")
    write_excel_sheets(s4_path, s4_sheets)
    save_csvs("Table_S4", s4_sheets)

    print("\n=== Step 5: Generating Table S5 (Cross-Tissue GO and Hallmark Pathway GSEA Profiling) ===")
    df_s5_go_raw = pd.read_csv(os.path.join(REPO_ROOT, "data/publication_input/go/04_tissue_statistics_concrete_terms_and_context.csv"))
    df_s5_go = df_s5_go_raw[['analysis_tissue', 'go_id', 'term_name', 'n_missions', 'n_analysis_units', 'mean_mission_nes', 'median_mission_nes', 'n_positive_missions', 'n_negative_missions', 'exact_signflip_p_two_sided', 'exact_signflip_FDR_all_cells', 'stouffer_meta_z', 'stouffer_p_two_sided', 'stouffer_FDR_all_cells']].rename(columns={
        'analysis_tissue': 'Tissue',
        'go_id': 'GO_ID',
        'term_name': 'GO_Term_Name',
        'mean_mission_nes': 'Mean_Mission_NES',
        'median_mission_nes': 'Median_Mission_NES',
        'exact_signflip_p_two_sided': 'Signflip_P_value',
        'exact_signflip_FDR_all_cells': 'Signflip_FDR',
        'stouffer_meta_z': 'Stouffer_Meta_Z',
        'stouffer_p_two_sided': 'Stouffer_P_value',
        'stouffer_FDR_all_cells': 'Stouffer_FDR'
    })

    df_s5_hallmark_raw = pd.read_csv(os.path.join(REPO_ROOT, "data/publication_input/hallmark/01_global_top_term_figure_selection.csv"))
    df_s5_hallmark = df_s5_hallmark_raw[['gene_set_name', 'term_display', 'direction', 'n_missions', 'mean_of_mission_mean_NES', 'median_of_mission_mean_NES', 'positive_missions', 'negative_missions']].rename(columns={
        'gene_set_name': 'Gene_Set_Name',
        'term_display': 'Pathway_Description',
        'direction': 'Direction',
        'mean_of_mission_mean_NES': 'Mean_Mission_NES',
        'median_of_mission_mean_NES': 'Median_Mission_NES',
        'positive_missions': 'Positive_Missions_Count',
        'negative_missions': 'Negative_Missions_Count'
    })

    s5_sheets = {
        "GO_Enrichment_Statistics": df_s5_go,
        "Hallmark_Top12_Pathways": df_s5_hallmark
    }
    s5_path = os.path.join(DESKTOP_DIR, "Table_S5_Functional_Pathway_and_Hallmark_GSEA.xlsx")
    write_excel_sheets(s5_path, s5_sheets)
    save_csvs("Table_S5", s5_sheets)

    print("\n=== Step 6: Generating Table S6 (Pathway Association and Cross-Tissue Fisher-z Network) ===")
    df_s6_fz_raw = pd.read_csv(os.path.join(REPO_ROOT, "data/publication_input/go/log2fc_spearman_20260902/04_overall_fisher_z_spearman_rho_matrix.csv"))
    df_s6_fz = df_s6_fz_raw.rename(columns={'term_key': 'Pathway_GO_Term'})

    df_s6_sp_raw = pd.read_csv(os.path.join(REPO_ROOT, "data/publication_input/go/log2fc_spearman_20260902/03_tissue_pairwise_spearman_log2fc_long.csv"))
    df_s6_sp = df_s6_sp_raw[['analysis_tissue', 'n_flight_samples', 'pathway_1_key', 'pathway_2_key', 'spearman_rho']].rename(columns={
        'analysis_tissue': 'Tissue',
        'n_flight_samples': 'Flight_Samples_Count',
        'pathway_1_key': 'Pathway_1',
        'pathway_2_key': 'Pathway_2',
        'spearman_rho': 'Spearman_Rho'
    })

    s6_sheets = {
        "Cross_Tissue_FisherZ_Matrix": df_s6_fz,
        "Per_Tissue_Spearman_Profiles": df_s6_sp
    }
    s6_path = os.path.join(DESKTOP_DIR, "Table_S6_Pathway_Association_and_FisherZ_Integration.xlsx")
    write_excel_sheets(s6_path, s6_sheets)
    save_csvs("Table_S6", s6_sheets)

    print("\n=== Step 7: Generating Table S7 (Double-Strand Break Repair Spaceflight Meta-Analysis) ===")
    df_s7_raw = pd.read_csv(os.path.join(REPO_ROOT, "data/publication_input/meta/08_essential_components_counts_log2FC_tissue_matrix.csv"))
    df_s7 = df_s7_raw[['Group', 'Pathway', 'EnsemblID', 'OfficialMouseSymbol', 'EntrezID', 'FlightExpressionLog2NormalizedCount', 'log2FoldChange', 'tissue_meta_p', 'padj', 'n_missions', 'n_analysis_units', 'expression_n_flight_samples']].rename(columns={
        'Group': 'Tissue',
        'OfficialMouseSymbol': 'Gene_Symbol',
        'FlightExpressionLog2NormalizedCount': 'Baseline_Log2_Expression',
        'log2FoldChange': 'Spaceflight_Log2FC',
        'tissue_meta_p': 'Meta_P_value',
        'padj': 'FDR_padj',
        'expression_n_flight_samples': 'Flight_Samples_Count'
    })
    s7_sheets = {"DSB_Meta_Analysis_26x29": df_s7}
    s7_path = os.path.join(DESKTOP_DIR, "Table_S7_DSB_Repair_Meta_Analysis_Matrix.xlsx")
    write_excel_sheets(s7_path, s7_sheets)
    save_csvs("Table_S7", s7_sheets)

    print("\n=== Step 8: Generating Table S8 (Single-Strand Break Repair Spaceflight Meta-Analysis) ===")
    df_s8_raw = pd.read_csv(os.path.join(REPO_ROOT, "release/tables/01_SSB_tissue_meta_log2FC_pvalue_matrix_without_Neil2.csv"))
    df_s8 = df_s8_raw[['Group', 'Pathway', 'EnsemblID', 'OfficialMouseSymbol', 'EntrezID', 'log2FoldChange', 'tissue_meta_p', 'padj', 'n_missions', 'n_analysis_units', 'expression_n_flight_samples']].rename(columns={
        'Group': 'Tissue',
        'OfficialMouseSymbol': 'Gene_Symbol',
        'log2FoldChange': 'Spaceflight_Log2FC',
        'tissue_meta_p': 'Meta_P_value',
        'padj': 'FDR_padj',
        'expression_n_flight_samples': 'Flight_Samples_Count'
    })
    s8_sheets = {"SSB_Meta_Analysis_26x45": df_s8}
    s8_path = os.path.join(DESKTOP_DIR, "Table_S8_SSB_Repair_Meta_Analysis_Matrix.xlsx")
    write_excel_sheets(s8_path, s8_sheets)
    save_csvs("Table_S8", s8_sheets)

    print("\n=== Step 9: Generating Master Consolidated Table S1-S4 (Teacher Citations) ===")
    master_sheets = {
        "S1_Cohort_Summary": df_s1_cohorts,
        "S1_Sample_Metadata_761": df_s1_samples,
        "S2_Thymus_Top30_DEGs": df_s2_thymus_deg,
        "S2_Kidney_Common12_DEGs": df_s2_kidney_deg,
        "S2_Up_Cohort_5_DEGs": df_s2_up_deg,
        "S2_Down_Cohort_11_DEGs": df_s2_down_deg,
        "S3_Pathway_Sizes": df_s3_sizes,
        "S3_Tissue_Detection": df_s3_tissue,
        "S4_ANOVA_Core_Genes": df_s4_gene_anova,
        "S4_Tissue_Gene_Mean_Log2": df_s4_tissue_mean
    }
    master_path = os.path.join(DESKTOP_DIR, "Table_S1_to_S4_Consolidated.xlsx")
    write_excel_sheets(master_path, master_sheets)

    print("\n=== Step 10: Manifest generation ===")
    manifest_rows = [
        {"Table_ID": "Table S1", "Title": "Animal Cohort Compilation and Biospecimen Metadata Atlas", "Teacher_Citation": "Fig. 1a-b, Table S1", "Sheets": "Cohort_Summary_Fig1b (52 cohorts × 12 cols), Sample_Metadata_Audit_761 (761 samples × 12 cols)", "File": "Table_S1_Cohort_and_Sample_Metadata.xlsx"},
        {"Table_ID": "Table S2", "Title": "Spaceflight-Induced Co-Directionally Regulated DDR Genes and Repertoire Overlap", "Teacher_Citation": "Fig. 3a, Fig. S2a, Fig. S3ab, Fig. S4a, Table S2", "Sheets": "Thymus_Top30_DEGs (30), Thymus_Overlap (893), Kidney_DEGs (12), Kidney_Overlap (893), Up_DEGs (5), Up_Overlap (893), Down_DEGs (11), Down_Overlap (893)", "File": "Table_S2_DDR_Overlap_and_Common_DEGs.xlsx"},
        {"Table_ID": "Table S3", "Title": "Seven DNA Damage Repair Pathways Gene Repertoire Sizes and Multi-Tissue Detection Audit", "Teacher_Citation": "Fig. S5, Table S3", "Sheets": "Pathway_Repertoire_Sizes (7 pathways × 2 cols), Tissue_Detection_Summary (182 rows × 12 cols), Sample_Pathway_Counts_761 (761 samples × 12 cols)", "File": "Table_S3_Seven_Pathways_Gene_Repertoires_and_Detection.xlsx"},
        {"Table_ID": "Table S4", "Title": "Baseline Expression Profiles and ANOVA Heterogeneity of Core Repair Genes Across Tissues", "Teacher_Citation": "Fig. 4b,c, Fig. 5c, Fig. 4/5 Legends, Table S4", "Sheets": "ANOVA_Core_Genes (26 genes × 8 cols), ANOVA_Pathways (7 pathways × 6 cols), Tissue_Gene_Mean_Log2 (676 rows × 9 cols), Sample_Log2_Values (9,360 rows × 8 cols)", "File": "Table_S4_Repair_Genes_Expression_and_ANOVA.xlsx"},
        {"Table_ID": "Table S5", "Title": "Cross-Tissue Gene Ontology Enrichment Profiling and Hallmark GSEA Overview", "Teacher_Citation": "Methods Section 2, Fig. 1c, Fig. S1", "Sheets": "GO_Enrichment_Statistics (390 rows: 26 tissues × 15 GO terms × 14 cols), Hallmark_Top12_Pathways (48 rows × 8 cols)", "File": "Table_S5_Functional_Pathway_and_Hallmark_GSEA.xlsx"},
        {"Table_ID": "Table S6", "Title": "Sample-Level Pathway Association and Cross-Tissue Fisher-z Network Integration", "Teacher_Citation": "Methods Section 3, Fig. 2, Fig. S6", "Sheets": "Cross_Tissue_FisherZ_Matrix (15 × 15 correlation matrix), Per_Tissue_Spearman_Profiles (2,730 pairwise correlations × 5 cols)", "File": "Table_S6_Pathway_Association_and_FisherZ_Integration.xlsx"},
        {"Table_ID": "Table S7", "Title": "Double-Strand Break (DSB) Repair Multi-Mission Spaceflight Meta-Analysis Matrix", "Teacher_Citation": "Methods Section 5, Fig. 4e", "Sheets": "DSB_Meta_Analysis_26x29 (754 rows: 26 tissues × 29 components × 12 cols)", "File": "Table_S7_DSB_Repair_Meta_Analysis_Matrix.xlsx"},
        {"Table_ID": "Table S8", "Title": "Single-Strand Break (SSB) Repair Multi-Mission Spaceflight Meta-Analysis Matrix", "Teacher_Citation": "Methods Section 6, Fig. 5f", "Sheets": "SSB_Meta_Analysis_26x45 (1,170 rows: 26 tissues × 45 components × 11 cols)", "File": "Table_S8_SSB_Repair_Meta_Analysis_Matrix.xlsx"},
    ]
    df_manifest = pd.DataFrame(manifest_rows)
    manifest_path = os.path.join(DESKTOP_DIR, "Table_S_Manifest.csv")
    df_manifest.to_csv(manifest_path, index=False)
    print(f"Manifest written: {manifest_path}")

if __name__ == "__main__":
    main()
